import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Create presentation
prs = Presentation()
# Set 16:9 ratio
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# Dark theme colors (from OLFI)
DARK_BG = RGBColor(10, 37, 37)     # #0A2525
EMERALD = RGBColor(16, 185, 129)   # #10B981
WARM_WHITE = RGBColor(225, 222, 209) # #E1DED1

def add_slide(title, subtitle=None, body_points=None, is_title_slide=False):
    layout = prs.slide_layouts[6] # Blank layout
    slide = prs.slides.add_slide(layout)
    
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    
    if is_title_slide:
        # Title
        txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(60)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        
        if subtitle:
            txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(11.33), Inches(1.5))
            tf2 = txBox2.text_frame
            p2 = tf2.paragraphs[0]
            p2.text = subtitle
            p2.font.size = Pt(28)
            p2.font.color.rgb = WARM_WHITE
            
    else:
        # Top accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(0.8), Inches(11.33), Inches(0.05))
        line.fill.solid()
        line.fill.fore_color.rgb = EMERALD
        line.line.fill.background()
        
        # Title
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11.33), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        
        current_top = 2.0
        
        if subtitle:
            txBox_sub = slide.shapes.add_textbox(Inches(1), Inches(current_top), Inches(11.33), Inches(0.8))
            tf_sub = txBox_sub.text_frame
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(22)
            p_sub.font.color.rgb = WARM_WHITE
            current_top += 1.0
            
        if body_points:
            txBox_body = slide.shapes.add_textbox(Inches(1), Inches(current_top), Inches(11.33), Inches(4.5))
            tf_body = txBox_body.text_frame
            for point in body_points:
                p_body = tf_body.add_paragraph()
                p_body.text = "• " + point
                p_body.font.size = Pt(18)
                p_body.font.color.rgb = WARM_WHITE
                p_body.space_after = Pt(14)
                
    return slide

# Build slides based on BuyOut_Pitch_V0.1.pptx strategy logic + Expansions

slides_data = [
    {
        "title": "BuyOut",
        "subtitle": "The infrastructure for digital debt refinancing in the GCC.",
        "is_title_slide": True
    },
    {
        "title": "The Opportunity",
        "subtitle": "Recent shifts in data availability, consumer behavior, and bank economics make a platform-led refinancing model viable today.",
        "body_points": [
            "Open banking & credit bureau APIs enable automated data gathering and verified profiling.",
            "Rising cost of consumer debt drives urgent consumer demand for savings and restructuring.",
            "Banks are under immense pressure to reduce Customer Acquisition Cost (CAC) for new financing.",
            "Consumers are willing to switch financial institutions for meaningful monthly savings."
        ]
    },
    {
        "title": "Product Introduction",
        "subtitle": "BuyOut delivers a simple, end-to-end refinancing experience that connects consumers, banks, and execution in a single flow.",
        "body_points": [
            "Consumers move from fragmented debt visibility to executed refinancing through a single, intuitive journey.",
            "The platform structures, prioritizes, and contextualizes decisions across all liabilities.",
            "Users connect accounts/inputs, BuyOut analyzes structure and cost, presents best consolidation options, and the bank executes.",
            "Focus is strictly on deep execution, moving beyond basic lead generation or rate comparison."
        ]
    },
    {
        "title": "The Product - Technical Architecture",
        "subtitle": "A robust, data-driven engine powered by real-time integrations and intelligent routing.",
        "body_points": [
            "Data Aggregation Layer: Secure connections to AECB and Open Banking APIs for holistic, real-time liability mapping.",
            "Intelligence Layer: Proprietary logic engines optimizing cost-savings, structured for single or multi-liability consolidation scenarios.",
            "Execution Layer: Direct API integrations with partner banks enabling seamless application data transfer and digital disbursement.",
            "Security & Compliance: Built from the ground up to adhere to CBUAE data residency, privacy, and local compliance standards."
        ]
    },
    {
        "title": "Product Uniqueness & Defensibility",
        "subtitle": "Defensibility is anchored in execution depth, multi-product consolidation, and deep bank-side integration.",
        "body_points": [
            "Execution Depth: Closing the loop with end-to-end flow rather than operating as a top-of-funnel aggregator.",
            "Multi-Product Consolidation: Simultaneously handling credit cards, auto loans, and personal loans in a single structuring event.",
            "Bank-Side Integration: Deep technical ties with lenders for direct application, underwriting data supply, and disbursement."
        ]
    },
    {
        "title": "Business Model",
        "subtitle": "An asset-light intermediary model monetizing successful refinancing without taking balance sheet risk.",
        "body_points": [
            "Revenue is generated through a phased monetization model aligned with platform maturity and partner value delivery.",
            "Phase 1: Lender referral and success fees generated per funded deal.",
            "Phase 2: Processing fees and advanced origination analytics for banking partners.",
            "Phase 3: Potential value-added services for consumers and financial institutions."
        ]
    },
    {
        "title": "Unit Economics & Market",
        "subtitle": "Simple unit economics driven by high-value events and near-zero marginal costs at scale.",
        "body_points": [
            "One successful refinance event generates significant commission value due to high ticket sizes.",
            "The marginal cost of processing approaches zero as bank API integrations mature and automate.",
            "The platform targets a massive, established refinanceable market where modest early penetration drives meaningful, sustainable revenue."
        ]
    },
    {
        "title": "Go-To-Market Strategy",
        "subtitle": "Built around clear product sequencing, a focused initial segment, and scalable acquisition channels.",
        "body_points": [
            "Balances high-intent digital demand generation with direct bank partnerships to drive volume.",
            "Scalable, digital-first acquisition channels ensure cost-effective top-of-funnel growth.",
            "Bank partnerships act as a core driver of conversion and unit economics, fundamentally de-risking the distribution model."
        ]
    },
    {
        "title": "GTM - Product Phasing",
        "subtitle": "Targeted, phased rollout to accelerate learning, maximize early conversion, and build volume.",
        "body_points": [
            "Phase 1 (M1-M6): Focus solely on credit cards and auto loans—areas with highest consumer pain points and lowest switching friction.",
            "Phase 2 (M7-M12): Expand to complex personal loans and holistic multi-product consolidation packages.",
            "Phase 3 (Year 2+): Rollout of advanced features, full market coverage, and integration of specialized/Sharia-compliant product mapping."
        ]
    },
    {
        "title": "GTM - Customer Persona",
        "subtitle": "The initial focus is salaried residents with multiple liabilities who are most motivated to refinance.",
        "body_points": [
            "High-Intent Segment: Users currently holding 2+ active credit facilities (e.g., credit cards, auto loans).",
            "Core Motivation: Immediate need to reduce monthly cash outflows, secure lower rates, or simplify fragmented debt management.",
            "Behavioral Profile: Digitally native, responsive to clear ROI (savings) messaging, and seeking frictionless, digital execution without branch visits."
        ]
    },
    {
        "title": "GTM - Bank & Creditor Partnerships",
        "subtitle": "Delivering pre-qualified, high-intent demand to create a compelling proposition for banks.",
        "body_points": [
            "What Banks Get: Fully vetted, pre-qualified applicants with structured liability data ready for underwriting.",
            "The Impact: Dramatically lowers customer acquisition costs (CAC) and increases refinancing conversion rates for partners.",
            "Strategic Alignment: De-risks the model for BuyOut while delivering high-LTV customers directly to partner institutions.",
            "Private Creditors: Specialist lenders benefit from an efficient, digitized, and highly targeted origination pipeline."
        ]
    },
    {
        "title": "Regulatory Positioning",
        "subtitle": "Operating as an intermediary platform, avoiding lending activity and credit origination entirely.",
        "body_points": [
            "Not a Lender: Zero balance sheet exposure or direct credit risk.",
            "Intermediary Role: Exclusively enables and executes refinancing transactions between consumers and licensed financial institutions.",
            "Compliance: Designed to operate securely within existing regulatory frameworks and open-finance mandates."
        ]
    },
    {
        "title": "Roadmap & Timeline",
        "subtitle": "Clear, measurable milestones from MVP to scale.",
        "body_points": [
            "Months 1-6: Launch MVP, integrate initial credit data APIs, and secure the first wave of bank partnerships for execution.",
            "Months 7-12: Expand product to handle multi-liability consolidation, aggressively scale user acquisition, and optimize the conversion funnel.",
            "Year 2: Deepen bank API integrations, introduce advanced intelligence features, and scale transaction volume to drive profitability.",
            "Year 3+: Evaluate regional expansion considerations and develop advanced data products."
        ]
    },
    {
        "title": "Execution & Credibility",
        "subtitle": "A team equipped to bridge technology, consumer finance, and bank operations.",
        "body_points": [
            "Deep domain expertise in UAE consumer finance, regulatory environments, and lending dynamics.",
            "Technical capability to build secure, scalable, API-first financial infrastructure.",
            "Strong local network required to execute strategic bank partnerships and navigate complex regulatory landscapes."
        ]
    },
    {
        "title": "The Ask",
        "subtitle": "Funding to launch MVP, secure partnerships, and capture early market share.",
        "body_points": [
            "Raising capital to fund critical technology development, initial customer acquisition, and operational setup.",
            "Funds directly unlock critical early milestones: core bank integrations and initial volume execution targets.",
            "Prepares the foundation for a larger Seed round within 12-18 months based on proven unit economics and volume."
        ]
    }
]

for sd in slides_data:
    add_slide(
        title=sd.get("title"),
        subtitle=sd.get("subtitle"),
        body_points=sd.get("body_points"),
        is_title_slide=sd.get("is_title_slide", False)
    )

prs.save("BuyOut_Pitch_Consultant_V1.pptx")
print("Presentation successfully saved as BuyOut_Pitch_Consultant_V1.pptx")
