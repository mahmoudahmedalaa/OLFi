import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# Premium Dark Theme
DARK_BG = RGBColor(10, 37, 37)     # #0A2525
CARD_BG = RGBColor(17, 51, 51)     # #113333
EMERALD = RGBColor(16, 185, 129)   # #10B981
WARM_WHITE = RGBColor(225, 222, 209) # #E1DED1
MUTED_TEXT = RGBColor(150, 160, 160)

def set_bg(slide):
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    # Top Emerald Accent Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
    line.fill.solid()
    line.fill.fore_color.rgb = EMERALD
    line.line.fill.background()

def add_header(slide, title, subtitle=None):
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.73), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    
    if subtitle:
        txBox_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.73), Inches(0.5))
        tf_sub = txBox_sub.text_frame
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.size = Pt(18)
        p_sub.font.color.rgb = WARM_WHITE

def create_title_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    
    # Large centered title box
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "BUYOUT"
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.alignment = PP_ALIGN.CENTER
    
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11.33), Inches(1.0))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "The infrastructure for digital debt refinancing in the GCC."
    p2.font.size = Pt(24)
    p2.font.color.rgb = WARM_WHITE
    p2.alignment = PP_ALIGN.CENTER
    return slide

def draw_cards(slide, cards_data, top=Inches(2.2), h=Inches(4.2)):
    count = len(cards_data)
    margin = Inches(0.8)
    gap = Inches(0.4)
    w = (prs.slide_width - (2*margin) - (gap*(count-1))) / count
    
    for i, card in enumerate(cards_data):
        left = margin + i*(w + gap)
        # Card Background
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = EMERALD
        shape.line.width = Pt(1.5)
        
        # Number/Icon (Optional)
        if "number" in card:
            tx_num = slide.shapes.add_textbox(left, top - Inches(0.3), w, Inches(0.8))
            p_num = tx_num.text_frame.paragraphs[0]
            p_num.text = card["number"]
            p_num.font.size = Pt(40)
            p_num.font.bold = True
            p_num.font.color.rgb = EMERALD
            p_num.alignment = PP_ALIGN.CENTER
        
        # Card Title
        tx_title = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.3), w - Inches(0.4), Inches(0.8))
        tf_title = tx_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = card["title"]
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = EMERALD
        p_title.alignment = PP_ALIGN.CENTER
        
        # Card Text
        tx_text = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(1.2), w - Inches(0.4), h - Inches(1.4))
        tf_text = tx_text.text_frame
        tf_text.word_wrap = True
        for point in card["points"]:
            p = tf_text.add_paragraph()
            p.text = point
            p.font.size = Pt(14)
            p.font.color.rgb = WARM_WHITE
            p.space_after = Pt(10)

def draw_timeline(slide, steps, top=Inches(3.0)):
    count = len(steps)
    margin = Inches(1.0)
    gap = Inches(0.2)
    w = (prs.slide_width - (2*margin) - (gap*(count-1))) / count
    h = Inches(1.5)
    
    for i, step in enumerate(steps):
        left = margin + i*(w + gap)
        
        # Chevron shape
        shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = EMERALD if i == 0 else CARD_BG
        shape.line.color.rgb = EMERALD
        shape.line.width = Pt(1)
        
        # Text inside chevron
        tx = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), w - Inches(0.5), h - Inches(0.4))
        tf = tx.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = step["phase"]
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = DARK_BG if i == 0 else EMERALD
        p.alignment = PP_ALIGN.CENTER
        
        # Text below
        tx_desc = slide.shapes.add_textbox(left, top + h + Inches(0.2), w, Inches(2.0))
        tf_desc = tx_desc.text_frame
        tf_desc.word_wrap = True
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = step["desc"]
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = WARM_WHITE

def draw_layers(slide, layers, top=Inches(2.5)):
    # 3 stacked horizontal bars representing tech stack
    count = len(layers)
    w = Inches(10)
    left = (prs.slide_width - w) / 2
    h = Inches(1.2)
    gap = Inches(0.3)
    
    for i, layer in enumerate(layers):
        curr_top = top + i*(h + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, curr_top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = EMERALD
        shape.line.width = Pt(1.5)
        
        # Title left side
        tx_title = slide.shapes.add_textbox(left + Inches(0.3), curr_top + Inches(0.2), Inches(3.0), h - Inches(0.4))
        tf_title = tx_title.text_frame
        tf_title.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_title = tf_title.paragraphs[0]
        p_title.text = layer["title"]
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = EMERALD
        
        # Desc right side
        tx_desc = slide.shapes.add_textbox(left + Inches(3.5), curr_top + Inches(0.2), w - Inches(3.8), h - Inches(0.4))
        tf_desc = tx_desc.text_frame
        tf_desc.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_desc.word_wrap = True
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = layer["desc"]
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = WARM_WHITE

def draw_split_metrics(slide, left_title, left_points, right_title, right_points):
    # Left Box
    w = Inches(5.5)
    h = Inches(4.5)
    top = Inches(2.2)
    left1 = Inches(0.8)
    
    shape1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left1, top, w, h)
    shape1.fill.solid()
    shape1.fill.fore_color.rgb = CARD_BG
    shape1.line.color.rgb = EMERALD
    
    tx1 = slide.shapes.add_textbox(left1 + Inches(0.3), top + Inches(0.3), w - Inches(0.6), Inches(0.5))
    tx1.text_frame.paragraphs[0].text = left_title
    tx1.text_frame.paragraphs[0].font.size = Pt(24)
    tx1.text_frame.paragraphs[0].font.color.rgb = EMERALD
    
    t1_body = slide.shapes.add_textbox(left1 + Inches(0.3), top + Inches(1.0), w - Inches(0.6), h - Inches(1.2))
    t1_body.text_frame.word_wrap = True
    for pt in left_points:
        p = t1_body.text_frame.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(16)
        p.font.color.rgb = WARM_WHITE
        p.space_after = Pt(12)
        
    # Right Box
    left2 = prs.slide_width - w - Inches(0.8)
    shape2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left2, top, w, h)
    shape2.fill.solid()
    shape2.fill.fore_color.rgb = CARD_BG
    shape2.line.color.rgb = EMERALD
    
    tx2 = slide.shapes.add_textbox(left2 + Inches(0.3), top + Inches(0.3), w - Inches(0.6), Inches(0.5))
    tx2.text_frame.paragraphs[0].text = right_title
    tx2.text_frame.paragraphs[0].font.size = Pt(24)
    tx2.text_frame.paragraphs[0].font.color.rgb = EMERALD
    
    t2_body = slide.shapes.add_textbox(left2 + Inches(0.3), top + Inches(1.0), w - Inches(0.6), h - Inches(1.2))
    t2_body.text_frame.word_wrap = True
    for pt in right_points:
        p = t2_body.text_frame.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(16)
        p.font.color.rgb = WARM_WHITE
        p.space_after = Pt(12)

# ----- SLIDE GENERATION -----

# 1. Title
create_title_slide()

# 2. The Opportunity (Cards)
s2 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s2)
add_header(s2, "The Opportunity", "Recent shifts in data availability, consumer behavior, and bank economics make a platform-led refinancing model viable today.")
draw_cards(s2, [
    {"number": "1", "title": "Open Banking", "points": ["AECB and Open Banking APIs enable automated data gathering.", "Verified profiling replaces manual paperwork."]},
    {"number": "2", "title": "Debt Costs", "points": ["Rising cost of consumer debt drives urgent consumer demand.", "Focus on immediate monthly savings."]},
    {"number": "3", "title": "Bank Pressure", "points": ["Banks are under immense pressure to reduce Customer Acquisition Cost (CAC).", "Need targeted, pre-qualified leads."]},
    {"number": "4", "title": "Switching", "points": ["Consumers are highly motivated to switch financial institutions for meaningful savings.", "Loyalty yields to cost-efficiency."]}
])

# 3. Product Uniqueness (Cards)
s3 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s3)
add_header(s3, "Product Uniqueness & Defensibility", "Defensibility is anchored in execution depth, multi-product consolidation, and deep bank-side integration.")
draw_cards(s3, [
    {"title": "Execution Depth", "points": ["Closing the loop with an end-to-end flow.", "Not just a top-of-funnel aggregator or rate comparison site.", "Full fulfillment tracking."]},
    {"title": "Multi-Product Consolidation", "points": ["Simultaneously handles credit cards, auto loans, and personal loans.", "Single structuring event for maximum consumer savings."]},
    {"title": "Bank-Side Integration", "points": ["Deep technical ties with lenders for direct application.", "Automated underwriting data supply and API disbursement."]}
])

# 4. Tech Architecture (Layers)
s4 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s4)
add_header(s4, "The Product: Technical Architecture", "A robust, data-driven engine powered by real-time integrations and intelligent routing.")
draw_layers(s4, [
    {"title": "1. Consumer Layer", "points": [], "desc": "React Native application featuring seamless KYC/OCR onboarding, intuitive dashboarding, and interactive consolidation scenarios."},
    {"title": "2. Intelligence Layer", "points": [], "desc": "Proprietary logic engine connected to AECB and Open Banking. Calculates structural savings, maps to Sharia-compliant products, and generates the BuyOut Score."},
    {"title": "3. Execution Layer", "points": [], "desc": "Direct API integrations with partner banks. Enables secure transfer of vetted underwriting data and facilitates digital disbursement."}
])

# 5. Business Model (Timeline/Phases)
s5 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s5)
add_header(s5, "Business Model Monetization", "An asset-light intermediary model monetizing successful refinancing without taking balance sheet risk.")
draw_timeline(s5, [
    {"phase": "Phase 1: Success Fees", "desc": "Lender referral and success fees generated per funded deal. Simple commission upon successful disbursement."},
    {"phase": "Phase 2: Processing", "desc": "Processing fees and advanced origination analytics packages sold to banking partners."},
    {"phase": "Phase 3: SaaS / Value-Add", "desc": "Premium value-added services for consumers and structured data products for financial institutions."}
])

# 6. GTM: Personas (Split)
s6 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s6)
add_header(s6, "GTM: Customer Target", "The initial focus is salaried residents with multiple liabilities who are most motivated to refinance.")
draw_split_metrics(s6, 
    "The Core Persona", 
    [
        "Target: Salaried UAE residents",
        "Profile: Holding 2+ active credit facilities (e.g., credit cards + auto loan).",
        "Motivation: Immediate need to reduce monthly cash outflows and simplify fragmented debt management.",
        "Behavior: Digitally native, responsive to clear ROI (savings) messaging."
    ],
    "Acquisition Strategy",
    [
        "High-intent digital demand generation (Google Search: 'debt consolidation').",
        "Performance social media targeting expats facing rate hikes.",
        "Employer channels positioning BuyOut as a financial wellness benefit.",
        "Frictionless execution replaces branch visits."
    ]
)

# 7. GTM: Bank Partnerships (Split)
s7 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s7)
add_header(s7, "GTM: Bank & Creditor Partnerships", "Delivering pre-qualified, high-intent demand to create a compelling proposition for banks.")
draw_split_metrics(s7, 
    "Value to Banks", 
    [
        "Fully vetted, pre-qualified applicants with structured liability data.",
        "Dramatically lowers Customer Acquisition Cost (CAC).",
        "Increases refinancing conversion rates and reduces processing time.",
        "Zero upfront balance sheet risk for BuyOut; pure origination pipeline."
    ],
    "Partnership Structure",
    [
        "Success fee structure (1.0% - 1.5% of funded value).",
        "Direct API integrations for data sharing and underwriting.",
        "Targeting tier 1 and tier 2 Islamic and conventional lenders.",
        "Specialist private creditors for highly targeted risk profiles."
    ]
)

# 8. Roadmap (Timeline)
s8 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s8)
add_header(s8, "Roadmap & Timeline", "Clear, measurable milestones from MVP to regional scale.")
draw_timeline(s8, [
    {"phase": "M1 - M6: MVP & Validation", "desc": "Launch Sandbox MVP, integrate AECB / Open Banking APIs, secure first 3 bank partnerships, acquire first 500 users."},
    {"phase": "M7 - M12: Scale & Consolidate", "desc": "Expand to multi-liability consolidation, scale user acquisition, achieve CAC efficiency, prove unit economics."},
    {"phase": "Year 2+: Expansion", "desc": "10,000+ users, deep underwriting API integrations, mature ML scoring, and preparation for KSA market entry."}
])

# 9. Regulatory Positioning (Cards)
s9 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s9)
add_header(s9, "Regulatory Positioning", "Operating securely within established frameworks.")
draw_cards(s9, [
    {"title": "Not a Lender", "points": ["Zero balance sheet exposure.", "No direct credit origination risk.", "Platform does not hold consumer funds."]},
    {"title": "Intermediary Role", "points": ["Exclusively enables and executes transactions.", "Bridges consumers and licensed financial institutions.", "Transparent matchmaking."]},
    {"title": "Compliance First", "points": ["Designed to operate securely within CBUAE open-finance mandates.", "DIFC / Sandbox regulatory pathways identified."]}
])

prs.save("BuyOut_Pitch_Consultant_V2.pptx")
print("V2 Saved.")
